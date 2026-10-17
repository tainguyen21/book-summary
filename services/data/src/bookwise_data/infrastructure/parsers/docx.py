"""DOCX parser adapter."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from bookwise_data.domain.source import (
    EmptySourceError,
    InvalidSourceError,
    ParsedSection,
    ParsedSource,
    PermanentProcessingError,
    SourceFormat,
    SourceLimitError,
)
from bookwise_data.infrastructure.parsers.limits import (
    MAX_DOCX_PARAGRAPHS,
    validate_archive_limits,
)
from bookwise_data.infrastructure.parsers.spans import bounded_spans


def parse_docx(source_content: bytes) -> ParsedSource:
    """Group paragraphs into sections that start at each heading."""

    try:
        validate_archive_limits(source_content, "docx")
        document = Document(BytesIO(source_content))
    except PermanentProcessingError:
        raise
    except Exception as error:
        raise InvalidSourceError(
            "invalid_docx_document",
            "The DOCX document could not be parsed.",
        ) from error

    if len(document.paragraphs) > MAX_DOCX_PARAGRAPHS:
        raise SourceLimitError(
            "docx_paragraph_limit",
            "The DOCX document exceeds the supported paragraph limit.",
        )

    groups: list[list[tuple[int, str, str]]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name
        if not groups or _node_type(style) == "heading":
            groups.append([])
        groups[-1].append((index, style, text))
    sections = tuple(_section(group) for group in groups)

    if not sections:
        raise EmptySourceError(
            "empty_docx_document",
            "The DOCX document contains no readable text.",
        )

    return ParsedSource(source_format=SourceFormat.DOCX, sections=sections)


def _section(paragraphs: list[tuple[int, str, str]]) -> ParsedSection:
    """A heading (if any) and the paragraphs up to the next heading."""

    first_index, first_style, first_text = paragraphs[0]
    return ParsedSection(
        node_type="section",
        title=first_text if _node_type(first_style) == "heading" else None,
        text="\n\n".join(text for _, _, text in paragraphs),
        location={"paragraph_number": first_index, "style": first_style},
        spans=tuple(
            span
            for index, style, text in paragraphs
            for span in bounded_spans(text, {"paragraph_number": index, "style": style})
        ),
    )


def _node_type(style_name: str) -> str:
    return "heading" if style_name.lower().startswith("heading") else "paragraph"
